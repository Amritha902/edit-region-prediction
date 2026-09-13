"""Build the submission documents as .docx — files that can be handed in."""
import json, sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

HERE = Path(__file__).resolve().parent
INK = RGBColor(0x1B, 0x1A, 0x17)
MUTE = RGBColor(0x5A, 0x58, 0x52)

def base_doc():
    d = Document()
    for name, size in (("Normal", 11), ):
        st = d.styles[name]; st.font.name = "Cambria"; st.font.size = Pt(size)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria")
    for lvl, size in ((1, 15), (2, 12.5), (3, 11.5)):
        st = d.styles[f"Heading {lvl}"]
        st.font.name = "Cambria"; st.font.size = Pt(size)
        st.font.color.rgb = INK; st.font.bold = True; st.font.italic = False
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.1)
    return d

def para(d, text, size=11, italic=False, bold=False, align=None, space=6, color=None):
    p = d.add_paragraph(); r = p.add_run(text)
    r.font.size = Pt(size); r.italic = italic; r.bold = bold
    if color is not None: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def rich(d, parts, size=11, space=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = d.add_paragraph()
    for t, b in parts:
        r = p.add_run(t); r.font.size = Pt(size); r.bold = b
    p.paragraph_format.space_after = Pt(space); p.paragraph_format.alignment = align
    return p

def set_cell(cell, text, bold=False, size=9.5, align="right"):
    cell.text = ""
    p = cell.paragraphs[0]; r = p.add_run(str(text))
    r.font.size = Pt(size); r.bold = bold; r.font.name = "Cambria"
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    p.alignment = {"right": WD_ALIGN_PARAGRAPH.RIGHT, "left": WD_ALIGN_PARAGRAPH.LEFT,
                   "center": WD_ALIGN_PARAGRAPH.CENTER}[align]

def booktabs(tbl):
    """Top and bottom rules, a rule under the header, nothing else."""
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "12"); el.set(qn("w:color"), "1B1A17"); borders.append(el)
    for edge in ("left", "right", "insideV"):
        el = OxmlElement(f"w:{edge}"); el.set(qn("w:val"), "none"); borders.append(el)
    el = OxmlElement("w:insideH"); el.set(qn("w:val"), "none"); borders.append(el)
    tblPr.append(borders)
    hdr = tbl.rows[0]
    for c in hdr.cells:
        tcPr = c._tc.get_or_add_tcPr(); b = OxmlElement("w:tcBorders")
        bo = OxmlElement("w:bottom"); bo.set(qn("w:val"), "single")
        bo.set(qn("w:sz"), "6"); bo.set(qn("w:color"), "1B1A17"); b.append(bo)
        tcPr.append(b)

def table(d, caption, header, rows, widths=None, bold_rows=(), align=None):
    cp = d.add_paragraph(); r = cp.add_run(caption)
    r.font.size = Pt(9.5); r.italic = True; r.font.color.rgb = MUTE
    cp.paragraph_format.space_after = Pt(3)
    t = d.add_table(rows=1, cols=len(header)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    align = align or (["left"] + ["right"] * (len(header) - 1))
    for i, h in enumerate(header):
        set_cell(t.rows[0].cells[i], h, bold=True, size=9.5, align=align[i])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, v in enumerate(row):
            set_cell(cells[i], v, bold=(ri in bold_rows), size=9.5, align=align[i])
    booktabs(t)
    if widths:
        for ri in range(len(t.rows)):
            for ci, w in enumerate(widths):
                t.rows[ri].cells[ci].width = Inches(w)
    d.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def figure(d, path, caption, width=6.2):
    if not Path(path).exists():
        print("  !! missing figure", path); return
    d.add_picture(str(path), width=Inches(width))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp = d.add_paragraph(); r = cp.add_run(caption)
    r.font.size = Pt(9.5); r.italic = True; r.font.color.rgb = MUTE
    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT; cp.paragraph_format.space_after = Pt(10)

def titlepage(d, kind, title, subtitle):
    for _ in range(3): d.add_paragraph()
    para(d, kind, size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=26, color=MUTE)
    p = d.add_paragraph(); r = p.add_run(title)
    r.font.size = Pt(21); r.bold = True; r.font.name = "Cambria"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(10)
    para(d, subtitle, size=12.5, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space=40, color=MUTE)
    for line, sz, it in (("Amritha S    23BEC1368", 12, False),
                         ("Yugeshwaran P    23BEC1404", 12, False),
                         ("", 6, False),
                         ("Guide: Dr. Saranya M (54783)", 11, True),
                         ("School of Electronics Engineering", 11, True),
                         ("Vellore Institute of Technology, Chennai", 11, True),
                         ("", 6, False),
                         ("Foundations of Data Science — Review 1", 11, False),
                         ("14 September 2026", 11, False)):
        para(d, line, size=sz, italic=it, align=WD_ALIGN_PARAGRAPH.CENTER, space=3)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def toc(d, entries):
    p = d.add_paragraph(); r = p.add_run("CONTENTS")
    r.font.size = Pt(12); r.bold = True
    p.paragraph_format.space_after = Pt(10)
    for num, name, lvl in entries:
        pp = d.add_paragraph()
        pp.paragraph_format.left_indent = Inches(0.28 if lvl == 2 else 0)
        pp.paragraph_format.space_after = Pt(2)
        r1 = pp.add_run(f"{num}   "); r1.font.size = Pt(10.5); r1.bold = (lvl == 1)
        r2 = pp.add_run(name); r2.font.size = Pt(10.5); r2.bold = (lvl == 1)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

from docx.opc.constants import RELATIONSHIP_TYPE as RT

def hyperlink(paragraph, text, url, size=9.5, bold=False):
    """A real clickable hyperlink — python-docx has no API for this."""
    rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rPr = OxmlElement("w:rPr")
    for tag, attr, val in (("w:color", "w:val", "1155CC"),
                           ("w:u", "w:val", "single"),
                           ("w:sz", "w:val", str(int(size*2))),
                           ("w:rFonts", "w:ascii", "Cambria")):
        el = OxmlElement(tag); el.set(qn(attr), val)
        if tag == "w:rFonts": el.set(qn("w:hAnsi"), "Cambria")
        rPr.append(el)
    if bold: rPr.append(OxmlElement("w:b"))
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t); h.append(r); paragraph._p.append(h)
    return h

def link_cell(cell, text, url, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if url and str(url).startswith("http"):
        hyperlink(p, text, url, size=size)
    else:
        r = p.add_run(text); r.font.size = Pt(size); r.font.name = "Cambria"
    return p
