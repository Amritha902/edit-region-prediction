"""Geometric QA: bounds, margins, overlaps, text fit."""
import sys
from pptx import Presentation

EMU_IN = 914400.0
SW, SH, MARGIN = 13.333, 7.5, 0.5
prs = Presentation(sys.argv[1])
issues = []
inch = lambda v: None if v is None else v / EMU_IN


def txt_of(sh):
    return "\n".join(p.text for p in sh.text_frame.paragraphs) if sh.has_text_frame else ""


for idx, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        x, y, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
        if None in (x, y, w, h):
            continue
        t = txt_of(sh).strip()
        if sh.has_table:
            # addTable rowH is a MINIMUM: real height is the sum of row heights.
            th = sum(r.height for r in sh.table.rows) / EMU_IN
            boxes.append((x, y, w, th, "[table]"))
            if y + th > SH - 0.6:
                issues.append(f"s{idx}: TABLE runs to {y+th:.2f} (past the footer rule)")
            continue
        if x < -0.01 or y < -0.01 or x + w > SW + 0.01 or y + h > SH + 0.01:
            if t or (w < 3 and h < 3):
                issues.append(f"s{idx}: OFF-SLIDE {t[:32]!r}")
        if t and 0 <= x < MARGIN - 0.01:
            issues.append(f"s{idx}: LEFT MARGIN {x:.2f} {t[:32]!r}")
        if t and y + h > SH - 0.10:
            issues.append(f"s{idx}: BOTTOM EDGE {y+h:.2f} {t[:32]!r}")
        if t and sh.has_text_frame and w > 0.35 and h > 0.15:
            sizes = [r.font.size.pt for p in sh.text_frame.paragraphs for r in p.runs if r.font.size]
            fs = max(sizes) if sizes else 12.0
            per = max(1, int(w / (fs * 0.50 / 72.0)))
            avail = max(1, int(h / (fs * 1.32 / 72.0)))
            est = sum(max(1, -(-len(l) // per)) for l in t.split("\n"))
            if est > avail:
                issues.append(f"s{idx}: TEXT FIT ~{est} lines vs {avail} ({fs:.0f}pt {w:.2f}x{h:.2f}) {t[:34]!r}")
        if t:
            boxes.append((x, y, w, h, t))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            ax, ay, aw, ah, at = boxes[i]; bx, by, bw, bh, bt = boxes[j]
            ox = min(ax+aw, bx+bw) - max(ax, bx); oy = min(ay+ah, by+bh) - max(ay, by)
            if ox > 0.14 and oy > 0.14 and ox*oy > 0.10:
                issues.append(f"s{idx}: OVERLAP {ox*oy:.2f}sq {at[:20]!r} <> {bt[:20]!r}")

print(f"{len(prs.slides)} slides checked\n")
for i in issues:
    print(" ", i)
print(f"\n{len(issues)} issue(s)" if issues else "No geometry issues found.")
